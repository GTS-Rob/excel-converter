import streamlit as st
import pandas as pd
from openpyxl import load_workbook
import re
from io import BytesIO
from datetime import datetime

def slugify(text):
    return re.sub(r'[^\w\s-]', '', text).strip().replace(' ', '_')

def get_excel_options(file_bytes):
    """Scan the workbook safely for sheets and tables."""
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    options = []
    for sheet_name in wb.sheetnames:
        options.append({'type': 'Sheet', 'name': sheet_name, 'parent': sheet_name, 'ref': None})
        ws = wb[sheet_name]
        if hasattr(ws, "tables") and ws.tables:
            for table in list(ws.tables.values()):
                options.append({'type': 'Table', 'name': table.displayName, 'parent': sheet_name, 'ref': table.ref})
    return options, wb

st.set_page_config(page_title="Excel to Markdown/CSV Converter", page_icon="📊")

st.title("📊 Excel to Markdown/CSV Converter")
st.markdown("Upload an Excel file and convert sheets or tables to Markdown or CSV format")

# File upload
uploaded_file = st.file_uploader("Choose an Excel file", type=['xlsx', 'xls'])

if uploaded_file is not None:
    file_bytes = uploaded_file.read()
    
    try:
        options, wb = get_excel_options(file_bytes)
        
        st.success(f"✅ File loaded: {uploaded_file.name}")
        
        # Selection dropdown
        option_labels = [f"{opt['type']}: {opt['name']}" for opt in options]
        selected_label = st.selectbox("Select a Sheet or Table", option_labels)
        selected_index = option_labels.index(selected_label)
        selected = options[selected_index]
        
        # Load data
        if selected['type'] == 'Sheet':
            df = pd.read_excel(BytesIO(file_bytes), sheet_name=selected['name'])
        else:
            ws = wb[selected['parent']]
            data = [[cell.value for cell in row] for row in ws[selected['ref']]]
            df = pd.DataFrame(data[1:], columns=data[0])
        
        # Cleanup
        df = df.dropna(how='all', axis=0).dropna(how='all', axis=1).fillna('')
        
        # Show columns
        st.subheader("📋 Available Columns")
        columns_list = list(df.columns)
        
        col1, col2 = st.columns([1, 3])
        with col1:
            for i, col in enumerate(columns_list):
                st.text(f"[{i}] {col}")
        
        # Column selection
        col_input = st.text_input(
            "Filter columns (enter numbers separated by commas, or leave blank for all)",
            placeholder="e.g., 0, 2, 5"
        )
        
        if col_input:
            try:
                indices = [int(x.strip()) for x in col_input.split(',')]
                selected_cols = [columns_list[i] for i in indices]
                df = df[selected_cols]
                st.info(f"Selected {len(selected_cols)} columns")
            except:
                st.error("Invalid column numbers. Please use format: 0, 1, 2")
        
        # Export format
        export_format = st.radio("Export Format", ["Markdown", "CSV"])
        
        # Preview
        st.subheader("📄 Preview (first 10 rows)")
        st.dataframe(df.head(10))
        
        # Generate output
        ts = datetime.now().strftime("%H%M%S")
        filename_base = f"{slugify(selected['name'])}_{ts}"
        
        if export_format == "CSV":
            output = df.to_csv(index=False)
            filename = f"{filename_base}.csv"
            mime_type = "text/csv"
        else:
            output = df.to_markdown(index=False, tablefmt="github")
            filename = f"{filename_base}.md"
            mime_type = "text/markdown"
        
        # Download button
        st.download_button(
            label=f"⬇️ Download {export_format}",
            data=output,
            file_name=filename,
            mime=mime_type
        )
        
        # Copy to clipboard section
        st.subheader("📋 Copy to Clipboard")
        st.code(output, language="markdown" if export_format == "Markdown" else "text")
        
    except Exception as e:
        st.error(f"❌ Error: {str(e)}")

else:
    st.info("👆 Upload an Excel file to get started")
